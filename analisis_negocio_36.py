# ============================================
# ANÁLISIS DE NEGOCIO - Bot de Ventas
# (continúa después de tu código de lectura, 
# consolidación y limpieza ya hecho)
# ============================================
import matplotlib.pyplot as plt
import pandas as pd
import seaborn as sns

# Importar el módulo que construye `df_consolidado`
import ventas

# Obtener el DataFrame consolidado creado en `ventas.py`
df_consolidado = ventas.df_consolidado

# Estilo visual
sns.set_style('whitegrid')
sns.set_palette('muted')
plt.rcParams.update({'figure.dpi': 150, 'axes.titlesize': 14, 'axes.labelsize': 11})

# --------------------------------------------
# PREGUNTA 1: ¿Cuánto vendió cada categoría en total?
# (EJEMPLO RESUELTO)
# --------------------------------------------
ventas_categoria = df_consolidado.groupby('categoria')['precio_unitario'].sum()
print("\n--- Ventas por categoria (total) ---")
print(ventas_categoria)

ventas_categoria_sorted = ventas_categoria.sort_values(ascending=False)
plt.figure(figsize=(8, 5))
ax = sns.barplot(x=ventas_categoria_sorted.index, y=ventas_categoria_sorted.values, palette='viridis')
ax.set_title('Ventas por categoría')
ax.set_ylabel('Ventas totales ($)')
ax.set_xlabel('Categoría')
ax.ticklabel_format(style='plain', axis='y')
for p in ax.patches:
	ax.annotate(f"{int(p.get_height()):,}", (p.get_x() + p.get_width() / 2., p.get_height()),
				ha='center', va='bottom', fontsize=9, color='black', xytext=(0, 4), textcoords='offset points')
plt.tight_layout()
plt.savefig("grafico_categoria.png", dpi=200)
plt.close()


# --------------------------------------------
# PREGUNTA 2: ¿Qué porcentaje de las ventas representa 
# cada vendedor?
# --------------------------------------------
# --------------------------------------------
# PREGUNTA 2: ¿Qué porcentaje de las ventas representa 
# cada vendedor?
# --------------------------------------------
ventas_vendedor = df_consolidado.groupby('vendedor')['precio_unitario'].sum()
print("\n--- Ventas por vendedor (total) ---")
print(ventas_vendedor)

porcentajes_vendedor = ventas_vendedor / ventas_vendedor.sum() * 100
vendedor_sorted = ventas_vendedor.sort_values(ascending=False)
colores = sns.color_palette('pastel', len(vendedor_sorted))
explode = [0.08 if i == 0 else 0 for i in range(len(vendedor_sorted))]
plt.figure(figsize=(8, 6))
plt.pie(vendedor_sorted.values, labels=vendedor_sorted.index, autopct='%1.1f%%', startangle=140, colors=colores, explode=explode, wedgeprops={'edgecolor': 'white'})
plt.title('Porcentaje de ventas por vendedor')
plt.tight_layout()
plt.savefig("grafico_vendedor.png", dpi=200)
plt.close()



# --------------------------------------------
# PREGUNTA 3: ¿Cuál es el producto que más se vende?
# --------------------------------------------
# --------------------------------------------
# PREGUNTA 3: ¿Cuál es el producto que más se vende?
# --------------------------------------------
producto_counts = df_consolidado['producto'].value_counts()
print("\n--- Recuento de ventas por producto ---")
print(producto_counts)

producto_top = producto_counts.index[0] if len(producto_counts) > 0 else None
print(f"Producto más vendido: {producto_top}")



# --------------------------------------------
# PREGUNTA 4: ¿Cómo se distribuyen las ventas según 
# el método de pago?
# --------------------------------------------
# --------------------------------------------
# PREGUNTA 4: ¿Cómo se distribuyen las ventas según 
# el método de pago?
# --------------------------------------------
ventas_metodo = df_consolidado.groupby('metodo_pago')['precio_unitario'].sum()
print("\n--- Ventas por método de pago (total) ---")
print(ventas_metodo)

ventas_metodo_sorted = ventas_metodo.sort_values(ascending=False)
plt.figure(figsize=(8, 5))
ax = sns.barplot(x=ventas_metodo_sorted.index, y=ventas_metodo_sorted.values, palette='cubehelix')
ax.set_title('Ventas por método de pago')
ax.set_ylabel('Ventas totales ($)')
ax.set_xlabel('Método de pago')
ax.ticklabel_format(style='plain', axis='y')
for p in ax.patches:
	ax.annotate(f"{int(p.get_height()):,}", (p.get_x() + p.get_width() / 2., p.get_height()),
				ha='center', va='bottom', fontsize=9, color='black', xytext=(0, 4), textcoords='offset points')
plt.tight_layout()
plt.savefig("grafico_metodo_pago.png", dpi=200)
plt.close()


# --------------------------------------------
# PREGUNTA 5 (RETO): ¿Cuál es el día de la semana con más ventas?
# --------------------------------------------
try:
	df_consolidado['fecha'] = pd.to_datetime(df_consolidado['fecha'], errors='coerce')
	df_consolidado['dia_semana'] = df_consolidado['fecha'].dt.day_name()
	ventas_dia = df_consolidado.groupby('dia_semana')['precio_unitario'].sum()
	# Asegurar orden razonable de días
	dias_orden_en = ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday', 'Saturday', 'Sunday']
	ventas_dia = ventas_dia.reindex(dias_orden_en).fillna(0)

	# Mapear nombres a español
	en_to_es = {
		'Monday': 'Lunes',
		'Tuesday': 'Martes',
		'Wednesday': 'Miércoles',
		'Thursday': 'Jueves',
		'Friday': 'Viernes',
		'Saturday': 'Sábado',
		'Sunday': 'Domingo',
	}
	ventas_dia_es = ventas_dia.rename(index=en_to_es)
	ventas_dia_es = ventas_dia_es[[en_to_es[d] for d in dias_orden_en]]

	print("\n--- Ventas por día de la semana (total) ---")
	print(ventas_dia_es)

	# Gráfico atractivo con seaborn y etiquetas en español
	plt.figure(figsize=(9, 5))
	ax = sns.barplot(x=ventas_dia_es.index, y=ventas_dia_es.values, palette='magma')
	ax.set_title('Ventas por día de la semana')
	ax.set_ylabel('Ventas totales ($)')
	ax.set_xlabel('Día de la semana')
	for p in ax.patches:
		ax.annotate(f"{int(p.get_height()):,}", (p.get_x() + p.get_width() / 2., p.get_height()),
					ha='center', va='bottom', fontsize=9, color='black', xytext=(0, 4), textcoords='offset points')
	plt.xticks(rotation=25)
	plt.tight_layout()
	plt.savefig('grafico_dia_semana.png', dpi=200)
	plt.close()
except Exception as e:
	print('No se pudo calcular P5:', e)


# --------------------------------------------
# GENERAR INFORME WORD - Analisis_Bot_Ventas.docx (MEJORADO)
# --------------------------------------------
try:
	from docx import Document
	from docx.shared import Inches
	import os

	doc = Document()
	doc.add_heading('Análisis Bot Ventas', level=1)

	# P1
	doc.add_heading('Pregunta 1: Ventas por categoría', level=2)
	for cat, val in ventas_categoria.items():
		doc.add_paragraph(f"{cat}: {val:,.2f}")
	if os.path.exists('grafico_categoria.png'):
		doc.add_picture('grafico_categoria.png', width=Inches(6))

	# P2
	doc.add_heading('Pregunta 2: Ventas por vendedor (porcentaje)', level=2)
	# Mostrar top 10 vendedores
	for vend, val in ventas_vendedor.sort_values(ascending=False).head(10).items():
		pct = val / ventas_vendedor.sum() * 100 if ventas_vendedor.sum() else 0
		doc.add_paragraph(f"{vend}: {val:,.2f} ({pct:.1f}%)")
	if os.path.exists('grafico_vendedor.png'):
		doc.add_picture('grafico_vendedor.png', width=Inches(5))

	# P3
	doc.add_heading('Pregunta 3: Producto que más se vende', level=2)
	doc.add_paragraph(f"Producto más vendido: {producto_top}")
	doc.add_paragraph('Top productos (cantidad de ventas):')
	for prod, cnt in producto_counts.head(10).items():
		doc.add_paragraph(f"{prod}: {cnt}")

	# P4
	doc.add_heading('Pregunta 4: Ventas por método de pago', level=2)
	for m, val in ventas_metodo.items():
		doc.add_paragraph(f"{m}: {val:,.2f}")
	if os.path.exists('grafico_metodo_pago.png'):
		doc.add_picture('grafico_metodo_pago.png', width=Inches(6))

	# P5 (reto)
	try:
		if 'ventas_dia' in locals() and not ventas_dia.empty:
			doc.add_heading('P5 (Reto): Ventas por día de la semana', level=2)
			for d, val in ventas_dia.items():
				doc.add_paragraph(f"{d}: {val:,.2f}")
			if os.path.exists('grafico_dia_semana.png'):
				doc.add_picture('grafico_dia_semana.png', width=Inches(6))
			top_day = ventas_dia.idxmax()
			doc.add_paragraph(f"Día con más ventas: {top_day} ({ventas_dia.max():,.2f})")
	except Exception:
		pass

	# Conclusión mejorada
	doc.add_heading('Conclusión y recomendaciones', level=2)
	conclusion = (
		"1) Refuerce stock y promociones para la categoría con mayor facturación. "
		"2) Diseñe incentivos para los vendedores que generan mayor venta y capacite a los de menor rendimiento. "
		"3) Asegure la experiencia en los métodos de pago más utilizados y simplifique el proceso de pago."
	)
	doc.add_paragraph(conclusion)

	# Guardar con fallback si el archivo está en uso
	salida = 'Analisis_Bot_Ventas.docx'
	try:
		doc.save(salida)
		print('\nInforme Word guardado en:', salida)
	except PermissionError:
		salida2 = 'Analisis_Bot_Ventas_actualizado.docx'
		doc.save(salida2)
		print('\nArchivo en uso. Informe guardado como:', salida2)

except Exception as e:
	print('\nNo se pudo generar el informe Word automáticamente. Requiere python-docx. Error:', e)


# --------------------------------------------
# RETO OPCIONAL - Para quien termine las 4 preguntas
# PREGUNTA 5: ¿Cuál es el día de la semana con más ventas?
# --------------------------------------------
# Paso 1: investiguen pd.to_datetime() para convertir la columna 
# fecha a formato de fecha real
# Paso 2: investiguen .dt.day_name() para extraer el día de la semana
# Paso 3: agrupen por ese nuevo dato y sumen las ventas